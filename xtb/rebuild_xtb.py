#!/usr/bin/env python3
"""Rebuild XTB account values (hourly, every 15 minutes for the last 60 days) from xStation5 history exports.

No scraping: the "Cash Operations" sheet of XTB's history export is a full
ledger — every deposit, withdrawal, purchase ("OPEN BUY 10 @ 100.00"),
sale, dividend, tax and IKE transfer, timestamped to the millisecond — so
cash and holdings can be replayed for any moment and priced from public
data:

  prices  Yahoo Finance 15-minute bars for the last 60 days (instruments held
          or traded then), hourly bars before, daily closes where hourly is
          missing (XTB ticker suffix mapped to Yahoo's, splits from Yahoo)
  FX      Yahoo 15-minute / hourly <CCY>PLN=X, NBP table A mid as the fallback

Value per account = cash + Σ quantity × last price × FX, with positions in a
currency other than the account's taken at market −0.5% (XTB's conversion
fee, which is how xStation values them). Splits are applied so quantities are
in today's units, matching Yahoo's split-adjusted prices. The result is
typically within ~0.1% of the value xStation shows. Leveraged positions
(CFDs, e.g. index or crypto CFDs) only contribute their realised cash, not open P/L.

Profit is all-time: total value − net external deposits (Deposit − Withdrawal,
each converted to PLN when it happened). It covers realised and open P/L,
dividends, interest, taxes and fees. IKE / sub-account transfers are internal.

Past the export's "Date to" (the cutoff) the series keeps going to now with
holdings frozen and live prices (after_export = 1). Your bank sees every transfer
to/from XTB, so bank rows booked after the cutoff that don't match a ledger
Deposit/Withdrawal (same amount ±0.5%, booked 0–7 days later) are added as
pending PLN cash of the main PLN account (XTB_PENDING_ACCOUNT, default: the
PLN_… export) *and* to net deposits — value moves, profit doesn't. Trades
since the cutoff stay invisible until the next export.

Usage:
  rebuild_xtb.py [EXPORT_DIR] [--out PATH] [--bank-tx transactions.csv]

EXPORT_DIR (default $XTB_EXPORT_DIR) is searched recursively for
`<CCY|IKE>_<account>_<from>_<to>.xlsx`; drop each new full-history export in
there (older ones of the same account can stay — ops are deduped by ID).
--out defaults to $XTB_REBUILT_FILE, --bank-tx to transactions.csv next to it.
One row per hour (Europe/Warsaw), every 15 minutes of weekday 9:00–18:00 in
the last 60 days, plus a final row at now:
  captured_at;after_export;total_pln;net_deposits_pln;profit_pln;pending_pln;
  cash_<acct>;positions_<acct>;balance_<acct>;currency_<acct>;…
Also writes xtb_holdings.csv next to --out: what each account holds now,
one row per instrument plus a CASH row per account:
  account;kind;ticker;name;qty;price;currency;value_pln;cost_pln
kind is the export's file prefix (PLN/IKE/EUR/USD); cost_pln is the average
purchase cost of the shares still held, at the FX of each purchase.
Runs on cron (XTB_REBUILD_CRON_SCHEDULE). Write is atomic (tmp → replace).
"""
import argparse
import bisect
import csv
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
import warnings
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import NamedTuple
from zoneinfo import ZoneInfo

import openpyxl

# XTB's xlsx carries no default style; openpyxl warns on every file.
warnings.filterwarnings("ignore", "Workbook contains no default style", UserWarning)

TZ = ZoneInfo("Europe/Warsaw")
FILE_RE = re.compile(r"^(PLN|EUR|USD|IKE)_(\d+)_")
TRADE_RE = re.compile(r"\b(OPEN|CLOSE) BUY ([\d.]+)(?:/[\d.]+)? @ ([\d.]+)")
# Money crossing the XTB boundary. Everything else (IKE deposit, Subaccount
# transfer) moves between your own accounts.
EXTERNAL_TYPES = {"Deposit", "Withdrawal"}
# XTB exchange suffix → Yahoo suffix.
YAHOO_SUFFIX = {"DE": ".DE", "FR": ".PA", "UK": ".L", "PL": ".WA", "NL": ".AS", "US": ""}
# Instruments Yahoo lists under another code. A different-currency listing of
# the same fund is fine — prices go through FX either way.
# Add one when a run logs CHECK MAPPING, e.g. "ABCD.DE": "WXYZ.DE".
YAHOO_OVERRIDE: dict[str, str] = {}
# Currency XTB quotes an instrument in, by suffix (.UK lines vary — trust Yahoo).
XTB_QUOTE_CCY = {"DE": "EUR", "FR": "EUR", "NL": "EUR", "PL": "PLN", "US": "USD"}
# Bank rows that are transfers to/from XTB (category_rules.csv excludes the same
# ones from spending). Unmatched ones after the cutoff land here.
BANK_XTB_RE = re.compile(r"X-?\s*TRADE|\bXTB\b", re.IGNORECASE)
BANK_FIELDS = ("creditor_name", "debtor_name", "remittance_information")
PENDING_MATCH_DAYS = 7
PENDING_MATCH_TOL = 0.005  # the bank can book a few grosze off XTB's amount (5 000.00 vs 4 999.97)
# XTB's currency-conversion fee: buys convert at market +0.5%, sells at −0.5%,
# and open positions are valued at the sell rate (verified against the
# open/close conversion rates in the export and the IKE open-position value).
FX_FEE = 0.005
UA = {"User-Agent": "Mozilla/5.0"}
# Yahoo serves 15-minute bars for the last 60 days only. Finer bars exist but
# thinly traded ETFs have no trade in most 1–5 minute slots.
FINE_INTERVAL, FINE_MINUTES, FINE_DAYS = "15m", 15, 59
# Yahoo symbol → instrument name (from the chart meta), for xtb_holdings.csv.
NAMES: dict[str, str] = {}


class Op(NamedTuple):
    time: datetime        # UTC
    account: str
    ticker: str
    amount: float         # cash delta, account currency
    qty: float            # holding delta (0 for non-trades)
    price: float | None   # fill price, instrument currency
    kind: str             # XTB operation type


def log(msg: str) -> None:
    print(f"[xtb-rebuild] {msg}", file=sys.stderr)


def get_json(url: str, attempts: int = 4):
    """GET + parse JSON, retrying timeouts and 5xx/429 with backoff (4xx raise at once)."""
    for i in range(attempts):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=30) as r:
                return json.load(r)
        except urllib.error.HTTPError as e:
            if e.code != 429 and e.code < 500 or i == attempts - 1:
                raise
        except (TimeoutError, urllib.error.URLError, ConnectionError):
            if i == attempts - 1:
                raise
        time.sleep(2 ** i * 3)


# --- ledger -----------------------------------------------------------------

def load_ledger(export_dir: Path):
    """→ {account: currency}, {account: PLN|EUR|USD|IKE}, [Op] sorted, cutoff
    (UTC "Date to" of the export)."""
    accounts, kinds, ops, seen, cutoffs = {}, {}, [], set(), {}
    for path in sorted(export_dir.rglob("*.xlsx")):
        m = FILE_RE.match(path.name)
        if not m:
            continue
        kind, number = m.groups()
        # IKE's own number lives in the filename; the sheet header can be wrong.
        accounts[number] = "PLN" if kind == "IKE" else kind
        kinds[number] = kind
        ws = openpyxl.load_workbook(path, data_only=True)["Cash Operations"]
        for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
            if row[0] == "Date to (UTC)" and isinstance(row[1], datetime):
                to = row[1].replace(tzinfo=timezone.utc)
                cutoffs[number] = max(cutoffs.get(number, to), to)
        for row in ws.iter_rows(min_row=6, values_only=True):
            typ, _, ticker, _, ts, amount, op_id, comment = row[:8]
            if not isinstance(ts, datetime) or amount is None or (number, op_id) in seen:
                continue
            seen.add((number, op_id))
            qty, price = 0.0, None
            if typ in ("Stock purchase", "Stock sell"):
                t = TRADE_RE.search(comment or "")
                if not t:
                    log(f"unparsed trade comment {comment!r} ({path.name})")
                else:
                    qty = float(t.group(2)) * (1 if t.group(1) == "OPEN" else -1)
                    price = float(t.group(3))
            ops.append(Op(ts.replace(tzinfo=timezone.utc), number, ticker or "",
                          float(amount), qty, price, typ))
    ops.sort()
    # The series is only trustworthy up to the account exported least recently.
    return accounts, kinds, ops, min(cutoffs.values())


class Ledger:
    """Cash and holdings per account as of any moment (bisect over cumulative snapshots)."""

    def __init__(self, accounts, ops: list[Op]):
        self.accounts = accounts
        self.ops = ops
        self.times = [o.time for o in ops]
        self.snaps = []
        cash = defaultdict(float)
        qty = defaultdict(lambda: defaultdict(float))
        for o in ops:
            cash[o.account] += o.amount
            if o.qty:
                qty[o.account][o.ticker] += o.qty
            self.snaps.append({a: (cash[a], {t: v for t, v in qty[a].items() if abs(v) > 1e-9})
                               for a in accounts})

    def at(self, t: datetime):
        i = bisect.bisect_right(self.times, t)
        return self.snaps[i - 1] if i else {a: (0.0, {}) for a in self.accounts}


# --- market data ------------------------------------------------------------

class Series:
    """Prices known at points in time; value at t = last point at or before t."""

    def __init__(self, points: dict[datetime, float]):
        self.times = sorted(points)
        self.values = [points[k] for k in self.times]

    def at(self, t: datetime) -> float:
        i = bisect.bisect_right(self.times, t)
        if not i:
            raise KeyError(f"no price before {t}")
        return self.values[i - 1]


def yahoo_bars(symbol: str, query: str):
    res = get_json(f"https://query1.finance.yahoo.com/v8/finance/chart/{symbol}?{query}")["chart"]["result"][0]
    quote = res["indicators"]["quote"][0]["close"]
    splits = [(datetime.fromtimestamp(s["date"], timezone.utc), s["numerator"] / s["denominator"])
              for s in res.get("events", {}).get("splits", {}).values()]
    return res["meta"], [(ts, c) for ts, c in zip(res.get("timestamp", []), quote) if c is not None], splits


def yahoo_series(symbol: str, start: date, fine: bool = False):
    """→ (currency, Series, splits). Intraday bars count from their end; daily
    closes from the end of their local day, so intraday data wins wherever it
    exists, and with `fine` the 15-minute bars add points between the hours. Yahoo's prices are split-adjusted; `splits` is [(utc_time, ratio)]
    so quantities can be brought to the same units (see adjust_for_splits)."""
    p1 = int(datetime.combine(start - timedelta(days=10), datetime.min.time(), timezone.utc).timestamp())
    meta, daily, splits = yahoo_bars(
        symbol, f"period1={p1}&period2={int(datetime.now().timestamp())}&interval=1d&events=split")
    try:
        _, hourly, _ = yahoo_bars(symbol, "range=730d&interval=1h")
    except urllib.error.HTTPError as e:
        log(f"{symbol}: no hourly bars (HTTP {e.code}), daily closes only")
        hourly = []
    finer = []
    if fine:
        try:
            _, finer, _ = yahoo_bars(symbol, f"range={FINE_DAYS}d&interval={FINE_INTERVAL}")
        except urllib.error.HTTPError as e:
            log(f"{symbol}: no {FINE_INTERVAL} bars (HTTP {e.code}), hourly only")
    NAMES[symbol] = meta.get("longName") or meta.get("shortName") or symbol
    tz = ZoneInfo(meta["exchangeTimezoneName"])
    scale = 0.01 if meta["currency"] == "GBp" else 1.0
    ccy = "GBP" if meta["currency"] == "GBp" else meta["currency"]
    points = {}
    for ts, c in daily:
        d = datetime.fromtimestamp(ts, tz).date()
        points[datetime.combine(d, datetime.max.time(), tz).astimezone(timezone.utc)] = c * scale
    now = datetime.now(timezone.utc)
    for ts, c in hourly:
        points[min(datetime.fromtimestamp(ts + 3600, timezone.utc), now)] = c * scale
    for ts, c in finer:
        points[min(datetime.fromtimestamp(ts + FINE_MINUTES * 60, timezone.utc), now)] = c * scale
    return ccy, Series(points), splits


def nbp_points(code: str, start: date, end: date) -> dict[datetime, float]:
    """NBP table A mid, PLN per unit, effective from noon Warsaw. Ranges cap at 93 days."""
    points, s = {}, start - timedelta(days=10)
    while s <= end:
        e = min(s + timedelta(days=92), end)
        data = get_json(f"https://api.nbp.pl/api/exchangerates/rates/a/{code.lower()}/{s}/{e}/?format=json")
        for r in data["rates"]:
            d = date.fromisoformat(r["effectiveDate"])
            points[datetime(d.year, d.month, d.day, 12, tzinfo=TZ).astimezone(timezone.utc)] = r["mid"]
        s = e + timedelta(days=1)
    return points


def fx_series(code: str, start: date) -> Series:
    """<code>→PLN: Yahoo hourly, NBP filling everything before Yahoo's history starts."""
    _, yahoo, _ = yahoo_series(f"{code}PLN=X", start, fine=True)
    first = yahoo.times[0]
    nbp = {t: v for t, v in nbp_points(code, start, first.date()).items() if t < first}
    return Series({**nbp, **dict(zip(yahoo.times, yahoo.values))})


def yahoo_symbol(ticker: str) -> str:
    if ticker in YAHOO_OVERRIDE:
        return YAHOO_OVERRIDE[ticker]
    base, _, suffix = ticker.rpartition(".")
    if suffix not in YAHOO_SUFFIX:
        raise ValueError(f"no Yahoo mapping for {ticker}")
    return base + YAHOO_SUFFIX[suffix]


def fetch_prices(ops: list[Op], start: date):
    """→ {ticker: (currency, Series, splits)} for every instrument ever traded.
    CFDs (index, crypto, …) have no exchange suffix and never become holdings.
    Instruments held now or traded in the 15-minute window get 15-minute bars."""
    tickers = sorted({o.ticker for o in ops if o.qty and "." in o.ticker})
    since = datetime.now(timezone.utc) - timedelta(days=FINE_DAYS)
    held = defaultdict(float)
    for o in ops:
        held[o.ticker] += o.qty
    recent = {o.ticker for o in ops if o.time >= since} | {t for t, q in held.items() if q > 1e-6}
    return {t: yahoo_series(yahoo_symbol(t), start, fine=t in recent) for t in tickers}


def adjust_for_splits(ops: list[Op], prices) -> list[Op]:
    """Express every trade in today's share units, matching Yahoo's split-adjusted
    prices (a 10:1 split: bought 1 @ 400.00 → 10 @ 40.00)."""
    out = []
    for o in ops:
        if o.qty and o.ticker in prices:
            factor = 1.0
            for when, ratio in prices[o.ticker][2]:
                if when > o.time:
                    factor *= ratio
            o = o._replace(qty=o.qty * factor, price=o.price / factor)
        out.append(o)
    return out


# --- valuation --------------------------------------------------------------

class Valuer:
    def __init__(self, ledger: Ledger, prices, start: date):
        self.ledger = ledger
        self.prices = prices
        ccys = ({p[0] for p in prices.values()} | set(ledger.accounts.values())) - {"PLN"}
        self.fx = {c: fx_series(c, start) for c in ccys}
        for when, snap in zip(ledger.times, ledger.snaps):
            for acct, (_, holdings) in snap.items():
                for ticker, q in holdings.items():
                    if q < -1e-6:
                        log(f"negative holding {ticker} {q:+.4f} in {acct} at {when:%Y-%m-%d} — missing trade or split?")
        # Net external deposits in PLN, cumulative, at the FX of the day they moved.
        self.dep_times, self.dep_cum, total = [], [], 0.0
        for o in ledger.ops:
            if o.kind in EXTERNAL_TYPES:
                total += o.amount * self.to_pln(ledger.accounts[o.account], o.time)
                self.dep_times.append(o.time)
                self.dep_cum.append(total)

    def check_trades(self) -> None:
        """XTB's own fill prices vs Yahoo at the same moment, per ticker.
        A median far from 0% means a wrong symbol mapping."""
        gaps = defaultdict(list)
        for o in self.ledger.ops:
            if o.price is None or o.ticker not in self.prices:
                continue
            ccy, series, _ = self.prices[o.ticker]
            xtb_ccy = XTB_QUOTE_CCY.get(o.ticker.rpartition(".")[2], ccy)
            try:
                yahoo = series.at(o.time) * self.to_pln(ccy, o.time) / self.to_pln(xtb_ccy, o.time)
            except KeyError:
                continue
            gaps[o.ticker].append((yahoo / o.price - 1) * 100)
        for ticker, g in sorted(gaps.items()):
            g.sort()
            med = g[len(g) // 2]
            if abs(med) > 3:
                log(f"fill check {ticker} → {yahoo_symbol(ticker)}: Yahoo vs XTB fill median "
                    f"{med:+.2f}% over {len(g)} trades — CHECK MAPPING")

    def to_pln(self, ccy: str, t: datetime) -> float:
        return 1.0 if ccy == "PLN" else self.fx[ccy].at(t)

    def net_deposits(self, t: datetime) -> float:
        i = bisect.bisect_right(self.dep_times, t)
        return self.dep_cum[i - 1] if i else 0.0

    def at(self, t: datetime):
        """→ {account: (cash, positions, total)} in each account's currency."""
        out = {}
        for acct, (cash, holdings) in self.ledger.at(t).items():
            acct_ccy = self.ledger.accounts[acct]
            pos = 0.0
            for ticker, q in holdings.items():
                ccy, series, _ = self.prices[ticker]
                v = q * series.at(t) * self.to_pln(ccy, t) / self.to_pln(acct_ccy, t)
                # XTB values foreign-currency positions at the rate it would sell
                # at: market minus its conversion fee.
                pos += v * (1 - FX_FEE) if ccy != acct_ccy else v
            out[acct] = (cash, pos, cash + pos)
        return out


def pending_moves(ledger: Ledger, cutoff: datetime, tx_path: Path) -> list[tuple[datetime, float]]:
    """Bank↔XTB transfers the export doesn't have yet → [(time, PLN into XTB)].

    Only rows booked on/after the cutoff day count; one that matches a ledger
    Deposit/Withdrawal booked 0–7 days earlier is already in the replay."""
    if not tx_path.exists():
        log(f"{tx_path} missing — no pending deposits")
        return []
    external = [(o.time.astimezone(TZ).date(), o.amount) for o in ledger.ops
                if o.kind in EXTERNAL_TYPES and ledger.accounts[o.account] == "PLN"]
    cutoff_day = cutoff.astimezone(TZ).date()
    pending = []
    with tx_path.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f, delimiter=";"):
            if not any(BANK_XTB_RE.search(row.get(k) or "") for k in BANK_FIELDS):
                continue
            booked = date.fromisoformat(row["booking_date"][:10])
            if booked < cutoff_day:
                continue
            amount = float(row["amount"]) * (1 if row["credit_debit"] == "DBIT" else -1)
            match = next((e for e in external
                          if 0 <= (booked - e[0]).days <= PENDING_MATCH_DAYS
                          and abs(e[1] - amount) <= abs(amount) * PENDING_MATCH_TOL), None)
            if match:
                external.remove(match)
                continue
            when = max(datetime.combine(booked, datetime.min.time(), TZ), cutoff.astimezone(TZ))
            pending.append((when, amount))
            log(f"pending bank↔XTB move {booked} {amount:+.2f} PLN (not in the export yet)")
    return sorted(pending)


def rebuild(valuer: Valuer, cutoff: datetime, pending: list[tuple[datetime, float]], pending_account: str):
    accts = sorted(valuer.ledger.accounts)
    header = ["captured_at", "after_export", "total_pln", "net_deposits_pln", "profit_pln", "pending_pln"]
    for a in accts:
        header += [f"cash_{a}", f"positions_{a}", f"balance_{a}", f"currency_{a}"]
    now = datetime.now(TZ).replace(microsecond=0)
    cutoff = min(cutoff, now)
    t = valuer.ledger.times[0].astimezone(TZ).replace(minute=0, second=0, microsecond=0) + timedelta(hours=1)
    times = set()
    while t < now:
        times.add(t)
        t += timedelta(hours=1)
    # Recent weekday trading hours every 15 minutes (the fine bars' window).
    t = (now - timedelta(days=FINE_DAYS)).replace(hour=0, minute=0, second=0)
    while t < now:
        if t.weekday() < 5 and 9 <= t.hour < 18:
            times.add(t)
        t += timedelta(minutes=FINE_MINUTES)
    # Exact rows at the cutoff (where frozen holdings start) and at now.
    times |= {cutoff.astimezone(TZ).replace(microsecond=0), now}
    rows = []
    for t in sorted(times):
        v = valuer.at(t)
        pend = sum(a for when, a in pending if when <= t)
        total_pln = sum(v[a][2] * valuer.to_pln(valuer.ledger.accounts[a], t) for a in accts) + pend
        deposits = valuer.net_deposits(t) + pend
        row = [t.isoformat(), int(t > cutoff), f"{total_pln:.2f}", f"{deposits:.2f}",
               f"{total_pln - deposits:.2f}", f"{pend:.2f}"]
        for a in accts:
            cash, pos, total = v[a]
            extra = pend if a == pending_account else 0.0
            row += [f"{cash + extra:.2f}", f"{pos:.2f}", f"{total + extra:.2f}", valuer.ledger.accounts[a]]
        rows.append(row)
    return header, rows


def holdings_now(valuer: Valuer, kinds, pending: list[tuple[datetime, float]], pending_account: str):
    """Each account's instruments and cash at now, valued like rebuild()'s last row.
    Cost is the average cost of the shares still held: buys add their cash (in
    PLN at the day's FX), sells remove the sold fraction."""
    now = datetime.now(TZ).replace(microsecond=0)
    ledger = valuer.ledger
    cost, held = defaultdict(float), defaultdict(float)
    for o in ledger.ops:
        if not o.qty or o.ticker not in valuer.prices:
            continue
        key = (o.account, o.ticker)
        if o.qty > 0:
            cost[key] += -o.amount * valuer.to_pln(ledger.accounts[o.account], o.time)
        elif held[key] > 1e-9:
            cost[key] *= max(held[key] + o.qty, 0.0) / held[key]
        held[key] += o.qty
    header = ["account", "kind", "ticker", "name", "qty", "price", "currency", "value_pln", "cost_pln"]
    rows = []
    for acct, (cash, holdings) in sorted(ledger.at(now).items()):
        acct_ccy = ledger.accounts[acct]
        for ticker, q in sorted(holdings.items()):
            ccy, series, _ = valuer.prices[ticker]
            price = series.at(now)
            v = q * price * valuer.to_pln(ccy, now)
            if ccy != acct_ccy:
                v *= 1 - FX_FEE
            rows.append([acct, kinds[acct], ticker, NAMES.get(yahoo_symbol(ticker), ticker), f"{q:.4f}",
                         f"{price:.4f}", ccy, f"{v:.2f}", f"{cost[(acct, ticker)]:.2f}"])
        if acct == pending_account:
            cash += sum(a for when, a in pending if when <= now)
        cash_pln = cash * valuer.to_pln(acct_ccy, now)
        rows.append([acct, kinds[acct], "CASH", acct_ccy, f"{cash:.2f}", "1", acct_ccy,
                     f"{cash_pln:.2f}", f"{cash_pln:.2f}"])
    return header, rows


def write_atomic(path: Path, header, rows) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(header)
        w.writerows(rows)
    tmp.replace(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("export_dir", type=Path, nargs="?", default=os.environ.get("XTB_EXPORT_DIR"))
    ap.add_argument("--out", type=Path, default=os.environ.get("XTB_REBUILT_FILE"))
    ap.add_argument("--bank-tx", type=Path, help="bank transactions.csv (default: next to --out)")
    args = ap.parse_args()
    if not args.export_dir or not Path(args.export_dir).is_dir():
        log(f"export dir {args.export_dir!r} missing — set XTB_EXPORT_DIR or pass it")
        return 1
    if not args.out:
        log("no output path — set XTB_REBUILT_FILE or pass --out")
        return 1
    out = Path(args.out)

    accounts, kinds, ops, cutoff = load_ledger(Path(args.export_dir))
    if not ops:
        log(f"no ledger ops found under {args.export_dir}")
        return 1
    log(f"{len(ops)} ledger ops, {ops[0].time:%Y-%m-%d} → {ops[-1].time:%Y-%m-%d}, "
        f"export cutoff {cutoff.astimezone(TZ):%Y-%m-%d %H:%M}, accounts {accounts}")
    start = ops[0].time.date()
    prices = fetch_prices(ops, start)
    ledger = Ledger(accounts, adjust_for_splits(ops, prices))
    valuer = Valuer(ledger, prices, start)
    valuer.check_trades()
    pending = pending_moves(ledger, cutoff, args.bank_tx or out.parent / "transactions.csv")
    # Bank transfers land in the main PLN account until the next export.
    pending_account = os.environ.get("XTB_PENDING_ACCOUNT") or next(
        (a for a, k in sorted(kinds.items()) if k == "PLN"), "")
    header, rows = rebuild(valuer, cutoff, pending, pending_account)
    write_atomic(out, header, rows)
    write_atomic(out.parent / "xtb_holdings.csv", *holdings_now(valuer, kinds, pending, pending_account))
    log(f"wrote {len(rows)} rows → {out} (last {rows[-1][0]}: total {rows[-1][2]} PLN, "
        f"profit {rows[-1][4]} PLN, {sum(r[1] for r in rows)} rows past the export)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
